import json
from io import BytesIO

import httpx
import pytest
from docx import Document as WordDocument
from sqlalchemy import select

from app.agent.intents import study_intent
from app.core.errors import AppError
from app.models import Document, StudyArtifact
from tests.test_agent import call


def study_book_bytes(chapters=9, long=False):
    book = WordDocument()
    for chapter in range(1, chapters + 1):
        book.add_heading(f"第{chapter}章 学习主题{chapter}", level=1)
        for index in range(60 if long and chapter == 1 else 3):
            book.add_paragraph(
                f"知识点{chapter}-{index}：训练模型需要区分训练资料与验证资料，学习应结合定义、条件、实例和反思。"
                * (4 if long else 1)
            )
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


@pytest.fixture
def study_app(exam_app):
    app, client, course, _ = exam_app
    uploaded = client.post(
        f"/api/courses/{course}/documents",
        files={"file": ("机器学习教材.docx", study_book_bytes())},
        data={"document_type": "textbook"},
    )
    assert uploaded.status_code == 201, uploaded.text
    return app, client, course, uploaded.json()["id"]


def test_natural_request_reads_all_chapters_without_planner(study_app, monkeypatch):
    app, client, course, book = study_app

    def never(*args):
        raise AssertionError("Whole material study must not use top-k retrieval or planner routing")

    monkeypatch.setattr(app.state.agent.planner, "step", never)
    monkeypatch.setattr(app.state.retriever, "search", never)
    result = client.post(
        "/api/agent",
        json={
            "course_id": course,
            "document_ids": [book],
            "question": "对于机器学习这个本书，提取一下每张重点以及核心知识点为我出一份复习卷子",
        },
    )
    assert result.status_code == 200, result.text
    data = result.json()
    assert (
        data["trace"][0]["tool"] == "generate_study_pack" and data["trace"][0]["status"] == "ok"
    ), data
    assert data["finish_reason"] == "completed" and not data["insufficient_evidence"]
    artifact = client.get(f"/api/courses/{course}/artifacts/{data['artifacts'][0]['id']}").json()
    assert artifact["kind"] == "study_pack" and len(artifact["content"]["sections"]) == 9
    assert len(artifact["content"]["questions"]) == 9
    assert len({q["topic"] for q in artifact["content"]["questions"]}) == 9
    with app.state.sessions() as session:
        assert artifact["coverage"]["read_chunks"] == session.get(Document, book).chunk_count
    assert {s["document_id"] for s in artifact["sources"]} == {book}
    assert "演示" in artifact["warning"]
    url = f"/api/courses/{course}/artifacts/{artifact['id']}/download"
    assert "参考答案与解析" in client.get(url).text
    question_only = client.get(url + "?include_answers=false").text
    assert "以下为检索原文" not in question_only and "参考答案与解析" not in question_only
    practice = client.post(
        f"/api/courses/{course}/practice",
        json={
            "reference_type": "artifact",
            "reference_id": artifact["id"],
            "question_index": 0,
            "outcome": "wrong",
        },
    )
    assert practice.status_code == 201


def model_pack_content(data, call_index):
    task, evidence = data["task"], data["evidence"]
    if task["kind"] == "review":
        topic = task["topics"][0]
        return {
            "title": "章节复习",
            "sections": [
                {
                    "topic": topic,
                    "summary": "核心概念与适用条件。",
                    "key_points": ["区分训练与验证。"],
                    "pitfalls": ["避免将验证资料用于训练。"],
                    "citation_ids": [evidence["topics"][topic][0]],
                }
            ],
        }
    questions = []
    for topic, quota in evidence["question_blueprint"].items():
        for index in range(quota):
            questions.append(
                {
                    "topic": topic,
                    "type": task["question_type"],
                    "difficulty": task["difficulty"],
                    "stem": f"练习{call_index}-{index}：说明主题「{topic}」中训练与验证的区别。",
                    "answer": "训练用于学习，验证用于检验。",
                    "explanation": "应保留独立验证资料。",
                    "citation_ids": [evidence["topics"][topic][0]],
                }
            )
    return {"title": "复习卷", "questions": questions}


def test_api_reads_every_chunk_consolidates_and_batches_paper(study_app):
    app, client, course, _ = study_app
    uploaded = client.post(
        f"/api/courses/{course}/documents",
        files={"file": ("长篇.docx", study_book_bytes(long=True))},
    )
    assert uploaded.status_code == 201
    book = uploaded.json()["id"]
    app.state.generation.settings.study_mode = "api"
    app.state.generation.settings.study_window_chars = 4000
    seen = []

    def handler(request):
        data = json.loads(json.loads(request.content)["messages"][1]["content"])
        seen.append(data)
        return httpx.Response(
            200,
            json={
                "choices": [
                    {"message": {"content": json.dumps(model_pack_content(data, len(seen)))}}
                ],
                "usage": {"total_tokens": 10},
            },
        )

    with httpx.Client(transport=httpx.MockTransport(handler)) as mock:
        app.state.generation.provider.client = mock
        result = client.post(
            "/api/study", json={"course_id": course, "document_ids": [book], "question_count": 20}
        )
    assert result.status_code == 201, result.text
    artifact = result.json()
    assert len(artifact["content"]["sections"]) == 9 and len(artifact["content"]["questions"]) == 20
    assert any(d["evidence"].get("consolidate") for d in seen)
    map_calls = [
        d for d in seen if d["task"]["kind"] == "review" and not d["evidence"].get("consolidate")
    ]
    observed = {s["id"] for d in map_calls for s in d["evidence"]["sources"]}
    assert len(observed) == artifact["coverage"]["read_chunks"]
    assert artifact["usage"]["total_tokens"] == 10 * len(seen)
    paper_calls = [d for d in seen if d["task"]["kind"] == "mock_exam"]
    assert len(paper_calls) >= 4 and all(d["task"]["question_count"] <= 6 for d in paper_calls)
    assert len({q["topic"] for q in artifact["content"]["questions"]}) == 9


def test_study_scope_readiness_budget_and_summary_only(study_app):
    app, client, course, book = study_app
    result = client.post(
        "/api/study", json={"course_id": course, "document_ids": [book], "include_questions": False}
    )
    assert result.status_code == 201 and result.json()["content"]["questions"] == []
    assert len(result.json()["content"]["sections"]) == 9
    app.state.generation.settings.study_max_calls = 4
    limited = client.post("/api/study", json={"course_id": course, "document_ids": [book]})
    assert limited.status_code == 422 and limited.json()["code"] == "study_scope_too_large"
    app.state.generation.settings.study_max_calls = 120
    before = len(client.get(f"/api/courses/{course}/artifacts").json())
    assert (
        client.post(
            "/api/study", json={"course_id": course, "document_ids": [book], "question_count": 3}
        ).status_code
        == 422
    )
    assert (
        client.post(
            "/api/study", json={"course_id": course, "document_name": "没有这个文件"}
        ).status_code
        == 409
    )
    other = client.post("/api/courses", json={"name": "another"}).json()["id"]
    assert (
        client.post("/api/study", json={"course_id": other, "document_ids": [book]}).status_code
        == 400
    )
    with app.state.sessions() as session:
        session.get(Document, book).status = "cancelled"
        session.commit()
    result = client.post(
        "/api/agent", json={"course_id": course, "question": "根据这本书整理每章重点并出复习卷"}
    )
    assert "已取消" in result.json()["answer"] and "题号" not in result.json()["answer"]
    assert len(client.get(f"/api/courses/{course}/artifacts").json()) == before


def test_provider_error_is_not_misreported_as_no_sources(study_app):
    app, client, course, _ = study_app
    app.state.agent.settings.study_mode = "api"

    class BrokenPlanner:
        calls = 0

        def step(self, messages, tools):
            self.calls += 1
            if self.calls == 1:
                return call("search_knowledge", {"query": "训练资料"})
            raise AppError("agent_invalid_output", "模型返回格式不合法", 502)

    app.state.agent.planner = BrokenPlanner()
    response = client.post(
        "/api/agent", json={"course_id": course, "question": "训练资料是什么"}
    ).json()
    assert response["sources"] and response["finish_reason"] == "provider_error"
    assert "模型响应发生错误" in response["answer"] and "未找到可靠来源" not in response["answer"]


def test_failed_chapter_generation_saves_no_partial_artifact(study_app):
    app, client, course, book = study_app
    app.state.generation.settings.study_mode = "api"
    calls = []

    def handler(request):
        data = json.loads(json.loads(request.content)["messages"][1]["content"])
        calls.append(data)
        content = model_pack_content(data, len(calls))
        if len(calls) == 2:
            content["sections"][0]["citation_ids"] = ["S999999"]
        return httpx.Response(
            200, json={"choices": [{"message": {"content": json.dumps(content)}}]}
        )

    with httpx.Client(transport=httpx.MockTransport(handler)) as mock:
        app.state.generation.provider.client = mock
        assert (
            client.post(
                "/api/study", json={"course_id": course, "document_ids": [book]}
            ).status_code
            == 502
        )
    with app.state.sessions() as session:
        assert list(session.scalars(select(StudyArtifact))) == []


@pytest.mark.parametrize(
    "question", ["生成薄弱点复习资料", "解释2025期末第1题", "A* 是什么", "生成 A* 模拟卷 2道简答题"]
)
def test_ordinary_queries_keep_existing_routes(question):
    assert study_intent(question) is None


@pytest.mark.parametrize("format_index", [0, 1, 2])
def test_study_pack_supported_material_formats(exam_app, tmp_path, format_index):
    from scripts.make_fixtures import generate

    _, client, course, _ = exam_app
    path = generate(tmp_path / "formats")[format_index]
    uploaded = client.post(
        f"/api/courses/{course}/documents",
        files={"file": (path.name, path.read_bytes())},
        data={"document_type": "notes"},
    )
    assert uploaded.status_code == 201
    result = client.post(
        "/api/study",
        json={
            "course_id": course,
            "document_ids": [uploaded.json()["id"]],
            "include_questions": False,
        },
    )
    assert result.status_code == 201, result.text
    assert result.json()["content"]["sections"]
    assert {s["document_id"] for s in result.json()["sources"]} == {uploaded.json()["id"]}
