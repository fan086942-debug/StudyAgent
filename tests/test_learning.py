import json
from io import BytesIO

import pytest
from docx import Document as WordDocument
from fastapi.testclient import TestClient

from app.agent.tools import ToolContext
from app.core.errors import AppError
from app.llm.base import ModelAnswer
from app.main import create_app
from app.schemas.agent import AgentTurn
from app.schemas.chat import SearchRequest
from tests.test_agent import ScriptedPlanner, call
from tests.test_generation import upload_material


def questions(client, course):
    return client.post(f"/api/courses/{course}/questions/search", json={}).json()


def practice(client, course, q, outcome="wrong", **extra):
    return client.post(
        f"/api/courses/{course}/practice",
        json={
            "reference_type": "question",
            "reference_id": q["id"],
            "outcome": outcome,
            "answer": "我的答案",
            "note": "混淆评价函数",
            **extra,
        },
    )


def test_wrongbook_deduplicates_and_latest_self_assessment(exam_app):
    _, client, course, _ = exam_app
    q = questions(client, course)[0]
    first = practice(client, course, q)
    assert first.status_code == 201
    assert practice(client, course, q).json()["wrong_id"] == first.json()["wrong_id"]
    wrongs = client.get(f"/api/courses/{course}/wrongbook").json()
    assert len(wrongs) == 1 and wrongs[0]["wrong_count"] == 2
    assert practice(client, course, q, "correct").status_code == 201
    assert client.get(f"/api/courses/{course}/wrongbook").json() == []
    progress = client.get(f"/api/courses/{course}/progress").json()
    assert progress["attempts"] == 3 and progress["practiced_questions"] == 1
    assert progress["accuracy"] == 100 and progress["mastered_wrong"] == 1
    assert sum(day["count"] for day in progress["week"]) == 3
    assert practice(client, course, q).status_code == 201
    assert client.get(f"/api/courses/{course}/progress").json()["accuracy"] == 0


def test_learning_scope_guards_and_validation(exam_app):
    _, client, course, _ = exam_app
    q = questions(client, course)[0]
    wrong = practice(client, course, q).json()["wrong_id"]
    other = client.post("/api/courses", json={"name": "其他课程"}).json()["id"]
    assert practice(client, other, q).status_code == 404
    assert (
        client.patch(f"/api/courses/{other}/wrongbook/{wrong}", json={"mastered": True}).status_code
        == 404
    )
    assert client.get(f"/api/courses/{other}/wrongbook").json() == []
    assert practice(client, course, q, outcome="fake").status_code == 422
    assert practice(client, course, q, snapshot={"stem": "fake"}).status_code == 422
    assert client.get("/api/courses/missing/progress").status_code == 404


def test_snapshot_survives_reparse_and_corrects_from_wrongbook(exam_app):
    _, client, course, doc = exam_app
    q = questions(client, course)[0]
    wrong = practice(client, course, q).json()["wrong_id"]
    assert client.post(f"/api/documents/{doc}/extract-questions").status_code == 200
    result = client.post(
        f"/api/courses/{course}/practice",
        json={"reference_type": "wrongbook", "reference_id": wrong, "outcome": "correct"},
    )
    assert result.status_code == 201 and result.json()["wrong_id"] == wrong
    assert client.get(f"/api/courses/{course}/wrongbook?include_mastered=true").json()[0][
        "mastered"
    ]


def test_generated_practice_uses_server_snapshot(exam_app):
    _, client, course, _ = exam_app
    upload_material(client, course)
    artifact = client.post(
        "/api/generate",
        json={
            "course_id": course,
            "kind": "mock_exam",
            "topics": ["A*"],
            "question_count": 1,
        },
    ).json()
    result = client.post(
        f"/api/courses/{course}/practice",
        json={"reference_type": "artifact", "reference_id": artifact["id"], "outcome": "wrong"},
    )
    assert result.status_code == 201, result.text
    wrong = client.get(f"/api/courses/{course}/wrongbook").json()[0]
    assert wrong["snapshot"]["answer"] == artifact["content"]["questions"][0]["answer"]
    assert (
        client.post(
            f"/api/courses/{course}/practice",
            json={
                "reference_type": "artifact",
                "reference_id": artifact["id"],
                "outcome": "wrong",
                "question_index": 11,
            },
        ).status_code
        == 404
    )


def test_plans_task_toggle_and_cross_course_access(exam_app):
    _, client, course, _ = exam_app
    practice(client, course, questions(client, course)[0])
    plan = client.post(f"/api/courses/{course}/plans", json={"days": 3, "daily_minutes": 25}).json()
    assert len(plan["tasks"]) == 3 and plan["topics"]
    path = f"/api/courses/{course}/plans/{plan['id']}/tasks/{plan['tasks'][0]['id']}"
    assert client.patch(path, json={"completed": True}).status_code == 200
    assert client.get(f"/api/courses/{course}/progress").json()["completed_tasks"] == 1
    assert client.patch(path, json={"completed": False}).status_code == 200
    other = client.post("/api/courses", json={"name": "other"}).json()["id"]
    assert client.patch(path.replace(course, other), json={"completed": True}).status_code == 404
    assert client.post(f"/api/courses/{course}/plans", json={"days": 0}).status_code == 422
    assert client.post(f"/api/courses/{other}/plans", json={}).status_code == 409
    topics = ["A*", "BFS", "DFS", "启发式函数"]
    short = client.post(f"/api/courses/{course}/plans", json={"days": 2, "topics": topics}).json()
    assert {t for task in short["tasks"] for t in task["topics"]} == set(topics)


def test_learning_agent_routes_and_scope(exam_app):
    app, client, course, doc = exam_app
    practice(client, course, questions(client, course)[0])
    for query, tool in [
        ("查看我的学习进度", "get_learning_progress"),
        ("看看我的错题", "list_wrong_questions"),
        ("安排3天复习计划，每天20分钟", "create_study_plan"),
    ]:
        response = client.post("/api/agent", json={"course_id": course, "question": query})
        assert response.status_code == 200, response.text
        data = response.json()
        assert data["trace"][0]["tool"] == tool and not data["insufficient_evidence"]
    assert len(client.get(f"/api/courses/{course}/plans").json()[0]["tasks"]) == 3
    material = upload_material(client, course)
    scoped = app.state.learning.progress(
        SearchRequest(course_id=course, question="进度", document_ids=[material])
    )
    assert scoped["attempts"] == 0 and scoped["active_wrong"] == 0
    scoped = app.state.learning.progress(
        SearchRequest(course_id=course, question="进度", document_ids=[doc])
    )
    assert scoped["attempts"] == 1
    response = client.post(
        "/api/agent", json={"course_id": course, "question": "生成薄弱知识点复习资料"}
    )
    assert response.json()["trace"][0]["tool"] == "generate_weakness_review"
    assert response.json()["artifacts"]


def test_history_scope_and_restart_persistence(exam_app):
    app, client, course, _ = exam_app
    practice(client, course, questions(client, course)[0])
    client.post(f"/api/courses/{course}/plans", json={"days": 2})
    for kind in ("agent", "chat"):
        assert (
            client.post(
                f"/api/{kind}", json={"course_id": course, "question": "A* 是什么"}
            ).status_code
            == 200
        )
    history = client.get(f"/api/courses/{course}/conversations").json()
    assert len(history) == 2
    other = client.post("/api/courses", json={"name": "other"}).json()["id"]
    assert client.get(f"/api/courses/{other}/conversations/{history[0]['id']}").status_code == 404
    with TestClient(create_app(app.state.settings)) as restarted:
        assert len(restarted.get(f"/api/courses/{course}/wrongbook").json()) == 1
        assert len(restarted.get(f"/api/courses/{course}/plans").json()) == 1
        detail = restarted.get(f"/api/courses/{course}/conversations/{history[0]['id']}").json()
        assert detail["result"]["answer"] and detail["question"] == "A* 是什么"


def test_agent_mutations_require_explicit_task(exam_app):
    app, _, course, _ = exam_app
    context = ToolContext(SearchRequest(course_id=course, question="A* 是什么？"))
    with pytest.raises(AppError, match="尚未要求"):
        app.state.agent.registry.execute("create_study_plan", "{}", context)
    with pytest.raises(AppError, match="尚未要求"):
        app.state.agent.registry.execute("generate_review", json.dumps({"topics": ["A*"]}), context)
    assert app.state.learning.plans(course) == []
    assert app.state.generation.list(course) == []


def test_api_plan_cache_saves_once_and_returns_trusted_summary(exam_app):
    app, client, course, _ = exam_app
    app.state.agent.settings.study_mode = "api"
    app.state.agent.planner = ScriptedPlanner(
        [
            call("create_study_plan", {"days": 2}, "plan-1"),
            call("create_study_plan", {"days": 2}, "plan-2"),
            AgentTurn(
                final=ModelAnswer(
                    answer="伪造学习记录", citation_ids=[], insufficient_evidence=True
                )
            ),
        ]
    )
    result = client.post(
        "/api/agent", json={"course_id": course, "question": "安排2天复习计划"}
    ).json()
    assert [t["status"] for t in result["trace"]] == ["ok", "cached"]
    assert len(client.get(f"/api/courses/{course}/plans").json()) == 1
    assert "已保存 2 天" in result["answer"] and "伪造" not in result["answer"]
    assert not result["insufficient_evidence"] and result["plans"]


def test_generated_wrongbook_scope_follows_question_evidence(exam_app):
    app, client, course, _ = exam_app
    upload_material(client, course)
    word = WordDocument()
    word.add_paragraph("BFS 队列先进先出，按层次扩展节点。")
    word.add_paragraph("BFS 队列在所有边代价相同时可求得最短路径。")
    stream = BytesIO()
    word.save(stream)
    assert (
        client.post(
            f"/api/courses/{course}/documents",
            files={"file": ("BFS.docx", stream.getvalue())},
            data={"document_type": "notes"},
        ).status_code
        == 201
    )
    artifact = client.post(
        "/api/generate",
        json={
            "course_id": course,
            "kind": "mock_exam",
            "topics": ["A*", "BFS 队列"],
            "question_count": 2,
        },
    ).json()
    source_docs = {s["document_id"] for s in artifact["sources"]}
    assert len(source_docs) == 2
    for index, question in enumerate(artifact["content"]["questions"]):
        cited = {
            s["document_id"]
            for s in artifact["sources"]
            if s["citation_id"] in question["citation_ids"]
        }
        excluded = list(source_docs - cited)
        if excluded:
            assert (
                client.post(
                    f"/api/courses/{course}/practice",
                    json={
                        "reference_type": "artifact",
                        "reference_id": artifact["id"],
                        "question_index": index,
                        "outcome": "wrong",
                    },
                ).status_code
                == 201
            )
            request = SearchRequest(course_id=course, question="进度", document_ids=excluded)
            assert app.state.learning.progress(request)["attempts"] == 0
            assert app.state.learning.wrong_questions(request) == []
            break
    else:
        pytest.fail("fixture must include a question grounded in only one document")
