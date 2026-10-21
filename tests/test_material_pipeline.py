import json
from io import BytesIO

import httpx
import pytest
from docx import Document as WordDocument
from pptx import Presentation
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas
from sqlalchemy import select

from app.models import Document, StudyArtifact
from app.schemas.chat import SearchRequest
from app.schemas.generation import GenerationRequest


def material_bytes(format_name, topic, concept):
    pages = [
        f"2.1 导论\n2.7 {topic}\n2.8 其他内容",
        f"2.7 {topic}\n2.7.1 基本概念\n2.7.2 适用边界\n2.7.3 验证方法",
        f"{concept}\n{concept}包含准备、执行与复核三个方面。\n" + "教学平台水印\n" * 40,
        "适用边界\n必须先检查输入状态，再执行更新；状态不满足条件时应拒绝更新。",
        "2.8 其他内容\n本部分属于另一节，不能混进前一节的依据。",
    ]
    buffer = BytesIO()
    if format_name == "docx":
        book = WordDocument()
        for page in pages:
            book.add_paragraph(page)
        book.save(buffer)
    elif format_name == "pptx":
        book = Presentation()
        for text in pages:
            slide = book.slides.add_slide(book.slide_layouts[1])
            title, _, body = text.partition("\n")
            slide.shapes.title.text = title
            slide.placeholders[1].text = body
        book.save(buffer)
    else:
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        book = canvas.Canvas(buffer)
        for text in pages:
            book.setFont("STSong-Light", 12)
            for index, line in enumerate(text.splitlines()):
                book.drawString(40, 800-index*16, line)
            book.showPage()
        book.save()
    return buffer.getvalue()


@pytest.mark.parametrize("format_name", ["pdf", "pptx", "docx"])
@pytest.mark.parametrize("topic,concept", [
    ("缓存的淘汰与更新", "优先队列"),
    ("任务的分配与回收", "轮转分配"),
])
def test_generic_topic_body_question_and_scope(exam_app, format_name, topic, concept):
    app, client, course, _ = exam_app
    uploaded = client.post(
        f"/api/courses/{course}/documents",
        files={"file": (f"不同领域资料.{format_name}", material_bytes(format_name, topic, concept))},
        data={"document_type": "notes"},
    )
    assert uploaded.status_code == 201, uploaded.text
    identifier = uploaded.json()["id"]
    request = SearchRequest(course_id=course, document_ids=[identifier], question=topic)
    # Deliberately make vector similarity unavailable; keyword evidence must still work.
    app.state.retriever.settings.min_similarity = 1
    sources = app.state.retriever.topic(request).sources
    assert sources and any(concept in s.text for s in sources)
    assert any("必须先检查输入状态" in s.text for s in sources)
    assert all("不能混进前一节" not in s.text for s in sources)
    assert all(s.document_id == identifier for s in sources)
    if format_name != "docx":
        assert {s.page for s in sources} == {3, 4}
    question = request.model_copy(update={"question": f"{topic}章节中的{concept}包括几个方面？"})
    answer_sources = app.state.retriever.answer(question).sources
    assert concept in answer_sources[0].text
    assert "三个方面" in answer_sources[0].text
    evidence, topic_ids = app.state.generation._evidence(GenerationRequest(
        course_id=course, document_ids=[identifier], topics=[topic], kind="review"))
    assert any(concept in s.text for s in evidence)
    assert topic_ids[topic] == [s.citation_id for s in evidence]
    with app.state.sessions() as db:
        db.get(Document, identifier).status = "stale"
        db.commit()
    assert not app.state.retriever.topic(request).sources
    assert not app.state.retriever.answer(question).sources


def test_task_contract_repaired_once_then_validated_and_cached(exam_app):
    app, client, course, _ = exam_app
    uploaded = client.post(f"/api/courses/{course}/documents", files={
        "file": ("独立流程.docx", material_bytes("docx", "样本的读取与校验", "输入校验"))})
    document = uploaded.json()["id"]
    app.state.generation.settings.study_mode = "api"
    calls = []

    def handler(request):
        payload = json.loads(request.content)
        data = json.loads(payload["messages"][1]["content"])
        calls.append(payload)
        topic = data["task"]["topics"][0]
        section = dict(topic=topic, summary="从正文概括输入状态和更新步骤。",
                       key_points=["更新前检查输入状态。"], pitfalls=["不要跳过前置检查。"],
                       citation_ids=[data["evidence"]["topics"][topic][0]])
        if len(calls) == 1:
            section["topic"] = "模型自行改写了主题名称"
        return httpx.Response(200, json={
            "choices": [{"message": {"content": json.dumps(dict(title="总结", sections=[section]))}}],
            "usage": {"total_tokens": 10},
        })

    with httpx.Client(transport=httpx.MockTransport(handler)) as http:
        app.state.generation.provider.client = http
        payload = dict(course_id=course, document_ids=[document], include_questions=False)
        first = client.post("/api/study", json=payload)
        assert first.status_code == 201, first.text
        assert "逐字" in calls[1]["messages"][-1]["content"]
        artifact = first.json()
        assert artifact["usage"]["total_tokens"] == 10 * len(calls)
        before = len(calls)
        second = client.post("/api/study", json=payload)
        assert second.status_code == 201, second.text
        assert len(calls) == before
        assert second.json()["coverage"]["cached_summaries"] > 0
        assert second.json()["usage"] == {}
    with app.state.sessions() as db:
        assert len(list(db.scalars(select(StudyArtifact)))) == 2


def test_unnumbered_heading_uses_section_body_without_repeating_topic(exam_app):
    app, client, course, _ = exam_app
    book = WordDocument()
    book.add_heading("任务恢复机制", level=1)
    book.add_paragraph("先保存已完成步骤，再读取检查点继续；不能重做已经完成的工作。")
    book.add_heading("其他内容", level=1)
    book.add_paragraph("这段不属于恢复机制。")
    data = BytesIO()
    book.save(data)
    response = client.post(f"/api/courses/{course}/documents", files={"file": ("无编号标题.docx", data.getvalue())})
    assert response.status_code == 201
    sources = app.state.retriever.topic(SearchRequest(
        course_id=course, document_ids=[response.json()["id"]], question="任务恢复机制")).sources
    assert any("先保存已完成步骤" in source.text for source in sources)
    assert all("这段不属于" not in source.text for source in sources)
