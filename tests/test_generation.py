import json
from copy import deepcopy

import httpx
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select

from app.core.config import Settings
from app.main import create_app
from app.models import Document, StudyArtifact
from tests.test_exams import exam_bytes

MATERIAL = (
    "A* 使用评价函数 f(n)=g(n)+h(n)。\n"
    "A* 的 g(n) 表示从起点到当前节点的实际路径代价。\n"
    "A* 的 h(n) 表示从当前节点到目标的估计代价。\n"
    "BFS 使用先进先出的队列逐层扩展节点。\n"
    "BFS 在等边代价下可找到边数最少的路径。\n"
    "DFS 使用后进先出的栈沿分支深入搜索。"
)


def upload_material(client, course):
    response = client.post(
        f"/api/courses/{course}/documents",
        files={"file": ("知识笔记.docx", exam_bytes(MATERIAL))},
        data={"document_type": "notes"},
    )
    assert response.status_code == 201, response.text
    return response.json()["id"]


@pytest.fixture
def generation_app(exam_app):
    app, client, course, exam = exam_app
    material = upload_material(client, course)
    return app, client, course, material, exam


def body(course, **kwargs):
    return {"course_id": course, "kind": "review", "topics": ["A*"], **kwargs}


def test_topic_generation_includes_body_beyond_isolated_heading(generation_app):
    from io import BytesIO

    from docx import Document as WordDocument

    from app.schemas.generation import GenerationRequest

    app, client, course, _, _ = generation_app
    doc = WordDocument()
    doc.add_heading("第1章 层次聚类", level=1)
    explanation = "凝聚层次聚类从单点簇开始，合并距离最近的两个簇。" + "单链接取跨簇点对距离的最小值。" * 20
    doc.add_paragraph(explanation)
    buffer = BytesIO()
    doc.save(buffer)
    uploaded = client.post(
        f"/api/courses/{course}/documents",
        files={"file": ("有标题和正文的教材.docx", buffer.getvalue())},
        data={"document_type": "textbook"},
    )
    assert uploaded.status_code == 201, uploaded.text
    document = uploaded.json()["id"]
    app.state.retriever.settings.min_similarity = 0.99
    request = GenerationRequest(
        course_id=course, document_ids=[document], kind="mock_exam",
        topics=["层次聚类"], question_count=2,
    )
    sources, topic_ids = app.state.generation._evidence(request)
    assert any("单链接取跨簇点对距离的最小值" in source.text for source in sources)
    assert all(source.document_id == document for source in sources)
    assert topic_ids["层次聚类"] == [source.citation_id for source in sources]


def test_demo_review_history_scope_and_download(generation_app):
    app, client, course, material, exam = generation_app
    response = client.post("/api/generate", json=body(course, topics=[" A* ", "BFS", "A*"]))
    assert response.status_code == 201, response.text
    artifact = response.json()
    assert artifact["request"]["topics"] == ["A*", "BFS"]
    assert len(artifact["content"]["sections"]) == 2
    assert artifact["mode"] == "demo" and artifact["usage"] == {}
    assert {s["document_id"] for s in artifact["sources"]} == {material}
    path = f"/api/courses/{course}/artifacts"
    assert client.get(path).json()[0]["id"] == artifact["id"]
    assert client.get(path + "?offset=1").json() == []
    assert client.get(f"{path}/{artifact['id']}").json() == artifact
    download = client.get(f"{path}/{artifact['id']}/download")
    assert download.status_code == 200 and "attachment;" in download.headers["content-disposition"]
    assert "来源原文快照" in download.text and "第" in download.text
    other = client.post("/api/courses", json={"name": "other"}).json()["id"]
    assert client.get(f"/api/courses/{other}/artifacts/{artifact['id']}").status_code == 404
    assert (
        client.get(f"/api/courses/{other}/artifacts/{artifact['id']}/download").status_code == 404
    )
    assert (
        client.post("/api/generate", json=body(other, document_ids=[material])).status_code == 400
    )
    assert client.post("/api/generate", json=body(course, document_ids=[exam])).status_code == 409


def test_demo_mock_separate_answers_and_agent(generation_app):
    _, client, course, _, _ = generation_app
    response = client.post("/api/generate", json=body(course, kind="mock_exam", question_count=3))
    assert response.status_code == 201, response.text
    result = response.json()
    assert len(result["content"]["questions"]) == 3
    url = f"/api/courses/{course}/artifacts/{result['id']}/download"
    assert "参考答案与解析" in client.get(url).text
    question_only = client.get(url + "?include_answers=false").text
    assert "参考答案与解析" not in question_only and "来源原文快照" not in question_only
    for question in ["生成 A* 复习资料", "生成 A* 模拟卷 2道简答题"]:
        response = client.post("/api/agent", json={"course_id": course, "question": question})
        assert response.status_code == 200, response.text
        assert len(response.json()["artifacts"]) == 1
        assert response.json()["trace"][0]["status"] == "ok"


@pytest.mark.parametrize(
    "changes",
    [
        {"topics": []},
        {"topics": [" "]},
        {"topics": ["a"] * 9},
        {"question_count": 13},
        {"document_ids": []},
        {"surprise": True},
        {"kind": "mock_exam", "topics": ["A*", "BFS"], "question_count": 1},
        {"kind": "mock_exam", "question_type": "choice"},
    ],
)
def test_invalid_requests(generation_app, changes):
    _, client, course, _, _ = generation_app
    assert client.post("/api/generate", json=body(course, **changes)).status_code == 422
    assert client.get(f"/api/courses/{course}/artifacts").json() == []


def test_no_evidence_stale_docs_and_busy(generation_app):
    app, client, course, material, _ = generation_app
    app.state.generation.lock.acquire()
    try:
        assert client.post("/api/generate", json=body(course)).status_code == 409
    finally:
        app.state.generation.lock.release()
    with app.state.sessions() as session:
        session.get(Document, material).status = "stale"
        session.commit()
    assert (
        client.post("/api/generate", json=body(course)).json()["code"] == "generation_no_evidence"
    )
    assert client.post("/api/generate", json=body("nonexistent")).status_code == 404
    assert client.get(f"/api/courses/{course}/artifacts").json() == []


def model_content(task, evidence):
    topic = task["topics"][0]
    citation = evidence["topics"][topic][0]
    if task["kind"] == "review":
        return {
            "title": "A* 复习",
            "sections": [
                {
                    "topic": topic,
                    "summary": "评价函数为 g+h。",
                    "key_points": ["区分实际代价和估计代价。"],
                    "pitfalls": ["不要把启发值当作实际代价。"],
                    "citation_ids": [citation],
                }
            ],
        }
    return {
        "title": "A* 练习",
        "questions": [
            {
                "topic": topic,
                "type": "choice",
                "difficulty": task["difficulty"],
                "stem": "A* 的评价函数是什么？",
                "options": ["g+h", "g-h", "g*h", "g/h"],
                "answer": "A",
                "explanation": "评价函数由实际代价和估计代价相加。",
                "citation_ids": [citation],
            }
        ],
    }


@pytest.mark.parametrize("kind", ["review", "mock_exam"])
def test_api_provider_contract(generation_app, kind):
    app, client, course, _, _ = generation_app
    seen = []

    def handler(request):
        payload = json.loads(request.content)
        data = json.loads(payload["messages"][1]["content"])
        seen.append(data)
        return httpx.Response(
            200,
            json={
                "choices": [
                    {
                        "message": {
                            "content": json.dumps(model_content(data["task"], data["evidence"]))
                        }
                    }
                ],
                "usage": {"total_tokens": 99},
            },
        )

    with httpx.Client(transport=httpx.MockTransport(handler)) as mock:
        app.state.generation.settings.study_mode = "api"
        app.state.generation.provider.client = mock
        result = client.post(
            "/api/generate", json=body(course, kind=kind, question_type="choice", question_count=1)
        )
    assert result.status_code == 201, result.text
    assert result.json()["mode"] == "api" and result.json()["usage"] == {"total_tokens": 99}
    assert "course_id" not in seen[0]["task"] and seen[0]["output_schema"]
    assert all(set(s) == {"id", "text"} for s in seen[0]["evidence"]["sources"])


@pytest.mark.parametrize(
    "fault",
    [
        "fake_citation",
        "missing_topic",
        "wrong_count",
        "wrong_type",
        "duplicate",
        "insufficient",
        "page",
        "truncated",
        "invalid_json",
        "invalid_choice",
        "timeout",
    ],
)
def test_invalid_model_output_is_never_saved(generation_app, fault):
    app, client, course, _, _ = generation_app

    def handler(request):
        if fault == "timeout":
            raise httpx.ReadTimeout("test timeout", request=request)
        data = json.loads(json.loads(request.content)["messages"][1]["content"])
        content = model_content(data["task"], data["evidence"])
        q = content["questions"][0]
        if fault == "fake_citation":
            q["citation_ids"] = ["S999"]
        elif fault == "missing_topic":
            q["topic"] = "unrequested"
        elif fault == "wrong_count":
            content["questions"] = []
        elif fault == "wrong_type":
            q.update(type="short_answer", options=[], answer="g+h")
        elif fault == "duplicate":
            content["questions"].append(deepcopy(q))
        elif fault == "insufficient":
            content["insufficient_evidence"] = True
        elif fault == "page":
            q["explanation"] = "见第999页。"
        elif fault == "invalid_choice":
            q["answer"] = "E"
        return httpx.Response(
            200,
            json={
                "choices": [
                    {
                        "finish_reason": "length" if fault == "truncated" else "stop",
                        "message": {
                            "content": "not-json"
                            if fault == "invalid_json"
                            else json.dumps(content)
                        },
                    }
                ]
            },
        )

    with httpx.Client(transport=httpx.MockTransport(handler)) as mock:
        app.state.generation.settings.study_mode = "api"
        app.state.generation.provider.client = mock
        result = client.post(
            "/api/generate",
            json=body(
                course,
                kind="mock_exam",
                question_type="choice",
                question_count=2 if fault == "duplicate" else 1,
            ),
        )
    assert result.status_code in (409, 502, 504), result.text
    with app.state.sessions() as session:
        assert session.scalar(select(StudyArtifact)) is None
    assert not app.state.generation.lock.locked()


def test_restart_keeps_artifact_and_snapshot(tmp_path):
    settings = Settings(_env_file=None, data_dir=tmp_path, min_similarity=0)
    with TestClient(create_app(settings)) as client:
        course = client.post("/api/courses", json={"name": "AI"}).json()["id"]
        material = upload_material(client, course)
        original = client.post("/api/generate", json=body(course)).json()
        assert client.post(f"/api/documents/{material}/retry").status_code == 200
    with TestClient(create_app(settings)) as client:
        result = client.get(f"/api/courses/{course}/artifacts/{original['id']}")
        assert result.status_code == 200 and result.json() == original
