from io import BytesIO

import pytest
from docx import Document as WordDocument
from pptx import Presentation
from sqlalchemy import select

from app.agent.intents import relation_intent, study_intent
from app.models import Document, DocumentChunk
from app.services.material_learning import contains_point
from tests.test_exams import exam_bytes


def upload(client, course, name, content, kind="textbook"):
    result = client.post(
        f"/api/courses/{course}/documents",
        files={"file": (name, content)},
        data={"document_type": kind},
    )
    assert result.status_code == 201, result.text
    return result.json()["id"]


def material_bytes():
    doc = WordDocument()
    doc.add_paragraph("第 1 章 搜索算法")  # Unstyled, spaced chapter heading.
    doc.add_paragraph("A* 的评价函数为 f(n)=g(n)+h(n)。启发式函数估计到目标的代价。")
    doc.add_paragraph("启发式函数不应高估实际最优代价；一致性确保沿路径评价值不下降。")
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


@pytest.mark.parametrize(
    "question",
    [
        "提取这个PPT的知识点",
        "总结Word的重点内容",
        "摘取PDF教材核心知识点",
        "归纳上传课件的知识点",
        "总结这份资料的内容",
        "整理每张的核心知识点",
        "总结这个Word",
        "提炼教材的核心知识点，不要出题",
    ],
)
def test_generic_summary_intents(question):
    assert study_intent(question)["include_questions"] is False


@pytest.mark.parametrize(
    "question",
    [
        "学习进度",
        "生成薄弱点复习资料",
        "A*是什么",
        "请总结学习计划",
        "这份资料在哪里",
        "生成A*模拟卷2道题",
    ],
)
def test_non_material_intents(question):
    assert study_intent(question) is None


def test_bidirectional_links_scopes_and_deleted_sources(exam_app):
    app, client, course, exam = exam_app
    material = upload(client, course, "搜索算法.docx", material_bytes())
    path = f"/api/courses/{course}/material-question-links"
    forward = client.post(path, json={"material_id": material, "document_ids": [material]}).json()
    assert len(forward["links"]) == 2
    assert {s["document_id"] for s in forward["sources"]} == {material}
    assert {q["document_id"] for q in forward["questions"]} == {exam}
    assert any("启发式函数" in link["concepts"] for link in forward["links"])
    question = forward["questions"][0]
    reverse = client.post(
        path,
        json={
            "direction": "exam_to_material",
            "question_id": question["id"],
            "document_ids": [material],
            "exam_document_ids": [exam],
        },
    ).json()
    assert reverse["links"][0]["question_id"] == question["id"]
    assert reverse["sources"][0]["chapter"] == "第 1 章 搜索算法"
    with app.state.sessions() as session:
        chunks = list(
            session.scalars(select(DocumentChunk).where(DocumentChunk.document_id == material))
        )
    scoped = client.post(path, json={"chunk_ids": [chunks[0].id]}).json()
    assert {s["chunk_id"] for s in scoped["sources"]} <= {chunks[0].id}
    assert client.post(path, json={"chunk_ids": ["fake"]}).status_code == 409
    other = client.post("/api/courses", json={"name": "另一课程"}).json()["id"]
    assert (
        client.post(
            f"/api/courses/{other}/material-question-links", json={"material_id": material}
        ).status_code
        == 404
    )
    assert client.post(path, json={"exam_document_ids": ["foreign"]}).status_code == 400
    assert client.post(path, json={"question_id": "missing"}).status_code == 404
    assert client.delete(f"/api/documents/{material}").status_code == 200
    assert client.post(path, json={}).json()["links"] == []


def test_knowledge_summary_persistence_and_stale_sources(exam_app):
    app, client, course, _ = exam_app
    material = upload(client, course, "课件配套资料.docx", material_bytes())
    result = client.post(
        "/api/agent",
        json={"course_id": course, "document_ids": [material], "question": "总结这个Word的知识点"},
    )
    assert result.json()["artifacts"] and result.json()["trace"][0]["tool"] == "generate_study_pack"
    artifact_id = result.json()["artifacts"][0]["id"]
    artifact = client.get(f"/api/courses/{course}/artifacts/{artifact_id}").json()
    assert not artifact["content"]["questions"]
    path = f"/api/courses/{course}/material-knowledge?document_id={material}"
    knowledge = client.get(path).json()
    assert knowledge and knowledge[0]["key_points"] and knowledge[0]["sources"]
    with app.state.sessions() as session:
        doc = session.get(Document, material)
        doc.status = "stale"
        session.commit()
    assert client.get(path).json() == []
    with app.state.sessions() as session:
        doc = session.get(Document, material)
        doc.status = "ready"
        session.commit()
    assert client.get(path).json()
    assert client.delete(f"/api/documents/{material}").status_code == 200
    assert client.get(f"/api/courses/{course}/material-knowledge").json() == []
    # Historical artifacts keep their evidence snapshots while active knowledge excludes them.
    assert client.get(f"/api/courses/{course}/artifacts/{artifact_id}").status_code == 200


def test_agent_explicit_relations_and_no_exam_guidance(exam_app):
    _, client, course, _ = exam_app
    material = upload(client, course, "课程参考.docx", material_bytes())
    result = client.post(
        "/api/agent", json={"course_id": course, "question": "查找教材对应试卷"}
    ).json()
    assert result["trace"][0]["tool"] == "match_material_questions"
    assert result["learning"]["material_question_links"] and "找到" in result["answer"]
    assert relation_intent("试卷对应教材")["direction"] == "exam_to_material"
    assert relation_intent("教材对应试卷")["direction"] == "material_to_exam"
    response = client.post(
        f"/api/courses/{course}/material-question-links",
        json={"material_id": material, "exam_document_ids": [material]},
    ).json()
    assert not response["links"] and "整份资料学习包" in response["answer"]


def test_untagged_question_and_unrelated_content(exam_app):
    _, client, course, _ = exam_app
    upload(client, course, "课程参考.docx", material_bytes())
    exam = upload(
        client,
        course,
        "无知识点标签.docx",
        exam_bytes("1. 启发式函数估计到目标的代价是什么意思？\n2. 解释古代文学中的诗歌格律。"),
        "exam",
    )
    result = client.post(
        f"/api/courses/{course}/material-question-links", json={"exam_document_ids": [exam]}
    ).json()
    assert len(result["links"]) == 1
    assert result["links"][0]["method"] == "text_overlap"
    assert result["questions"][0]["number"] == "1"


@pytest.mark.parametrize("index", [0, 1, 2])
def test_links_supported_formats(exam_app, tmp_path, index):
    from scripts.make_fixtures import generate

    _, client, course, _ = exam_app
    path = generate(tmp_path / "formats")[index]
    material = upload(client, course, path.name, path.read_bytes())
    response = client.post(
        f"/api/courses/{course}/material-question-links", json={"material_id": material}
    )
    assert response.status_code == 200, response.text
    assert response.json()["links"]
    assert {s["file_type"] for s in response.json()["sources"]} == {path.suffix[1:]}


def test_long_slide_deck_keeps_all_content(exam_app):
    _, client, course, _ = exam_app
    deck = Presentation()
    for n in range(80):
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = f"概念{n}"
        slide.placeholders[1].text = f"第{n}节的算法定义与应用。测试该章节的完整知识内容。"
    stream = BytesIO()
    deck.save(stream)
    material = upload(client, course, "大课件.pptx", stream.getvalue(), "slides")
    result = client.post(
        "/api/study",
        json={"course_id": course, "document_ids": [material], "include_questions": False},
    )
    assert result.status_code == 201, result.text
    assert result.json()["coverage"]["read_chunks"] == 80
    assert len(result.json()["content"]["sections"]) <= 64
    assert "合并相邻小节" in result.json()["warning"]


def test_label_boundaries_do_not_match_unrelated_words():
    assert contains_point("The A* evaluation function", "A*")
    assert contains_point("LINEAR\nregression estimates a relationship", "Linear regression")
    assert not contains_point("scalar operations", "car")


def test_english_untagged_material_question_relation(exam_app):
    _, client, course, _ = exam_app
    doc = WordDocument()
    doc.add_paragraph("Linear regression minimizes squared errors to fit a relationship.")
    stream = BytesIO()
    doc.save(stream)
    material = upload(client, course, "Statistics.docx", stream.getvalue())
    exam = upload(
        client,
        course,
        "Statistics exam.docx",
        exam_bytes("1. Explain linear regression with squared errors."),
        "exam",
    )
    result = client.post(
        f"/api/courses/{course}/material-question-links",
        json={"material_id": material, "exam_document_ids": [exam]},
    ).json()
    assert result["links"] and result["links"][0]["method"] == "text_overlap"


def test_material_questions_are_generated_without_historical_paper():
    for question in ("根据PPT出5道题", "根据Word出卷", "基于上传资料生成练习题"):
        assert study_intent(question)["include_questions"] is True


def test_reused_chunk_id_does_not_revive_old_knowledge(exam_app):
    app, client, course, _ = exam_app
    material = upload(client, course, "版本变化.docx", material_bytes())
    artifact = client.post(
        "/api/study",
        json={"course_id": course, "document_ids": [material], "include_questions": False},
    ).json()
    assert client.get(f"/api/courses/{course}/material-knowledge").json()
    with app.state.sessions() as session:
        for chunk in session.scalars(
            select(DocumentChunk).where(DocumentChunk.document_id == material)
        ):
            chunk.text = "重建后的不同正文，原始总结已经不再有依据。"
        session.commit()
    assert client.get(f"/api/courses/{course}/material-knowledge").json() == []
    response = client.post(
        f"/api/courses/{course}/material-question-links",
        json={"artifact_id": artifact["id"], "chunk_ids": [artifact["sources"][0]["chunk_id"]]},
    )
    assert response.status_code == 409 and response.json()["code"] == "stale_knowledge"
