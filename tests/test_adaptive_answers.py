import json
from types import SimpleNamespace

import httpx
import pytest
from pptx import Presentation

from app.core.config import Settings
from app.llm.base import GenerationResult, ModelAnswer
from app.llm.compatible_provider import CompatibleLLM
from app.schemas.chat import SearchRequest, SearchResponse
from app.services.qa import QAService
from tests.test_agent import ScriptedPlanner, call


class AnswerLLM:
    def __init__(self, answer):
        self.answer = answer
        self.calls = []

    def solve(self, question, evidence):
        self.calls.append((question, evidence))
        value = self.answer(evidence) if callable(self.answer) else self.answer
        return GenerationResult(content=value, usage={"total_tokens": 11})

    def generate(self, *_args):
        raise AssertionError("The optional-reference solving prompt must be used")


def model_answer():
    return ModelAnswer(
        answer="模型推理解答：A* 的评价函数为 f(n)=g(n)+h(n)，选 A。",
        citation_ids=[], insufficient_evidence=False, answer_basis="model",
    )


def test_uploaded_question_answers_without_textbook(exam_app):
    app, client, course, _ = exam_app
    question = client.post(
        f"/api/courses/{course}/questions/search", json={"question_type": "choice"}
    ).json()[0]
    app.state.settings.study_mode = "api"
    provider = AnswerLLM(model_answer())
    app.state.agent.registry.llm = provider
    result = client.post(
        f"/api/questions/{question['id']}/analyze",
        json={"course_id": course, "question": "分析"},
    ).json()
    assert result["answer_basis"] == "model"
    assert not result["insufficient_evidence"]
    assert provider.calls[0][1] == []
    assert all(source["document_type"] == "exam" for source in result["sources"])
    assert result["sources"][0]["block_index"] == question["locations"][0]["block_index"]


@pytest.mark.parametrize("stream", [False, True])
def test_empty_course_direct_answer_and_history(exam_app, stream):
    app, client, _, _ = exam_app
    course = client.post("/api/courses", json={"name": "无需教材"}).json()["id"]
    app.state.settings.study_mode = "api"
    planner = ScriptedPlanner([call("answer_question", {"query": "A* 的评价函数是什么？"})])
    provider = AnswerLLM(model_answer())
    app.state.agent.planner = planner
    app.state.agent.registry.llm = provider
    session = "a" * 32
    response = client.post(
        "/api/agent/stream" if stream else "/api/agent",
        json={"course_id": course, "question": "A* 的评价函数是什么？", "session_id": session},
    )
    assert response.status_code == 200, response.text
    if stream:
        events = [json.loads(line[6:]) for line in response.text.splitlines() if line.startswith("data: ")]
        result = next(event["data"] for event in events if event["type"] == "result")
    else:
        result = response.json()
    assert result["answer_basis"] == "model" and result["sources"] == []
    assert result["finish_reason"] == "completed" and not result["insufficient_evidence"]
    assert len(planner.messages) == 1  # A verified solver answer needs no second planning call.
    history = client.get(f"/api/agent/sessions/{session}", params={"course_id": course}).json()
    assert history["turns"][-1]["answer_basis"] == "model"
    assert history["turns"][-1]["answer"] == result["answer"]


def test_referenced_solution_keeps_actual_slide_only(exam_app, tmp_path):
    app, client, course, _ = exam_app
    slides = Presentation()
    slides.slides.add_slide(slides.slide_layouts[6])
    slide = slides.slides.add_slide(slides.slide_layouts[6])
    box = slide.shapes.add_textbox(0, 0, 7000000, 4000000)
    box.text_frame.text = "A* 的评价函数为 f(n)=g(n)+h(n)，其中 h(n) 是启发式估计。"
    path = tmp_path / "算法课件.pptx"
    slides.save(path)
    response = client.post(
        f"/api/courses/{course}/documents", files={"file": (path.name, path.read_bytes())},
        data={"document_type": "slides"},
    )
    assert response.status_code == 201, response.text
    app.state.settings.study_mode = "api"
    app.state.agent.planner = ScriptedPlanner([call("answer_question", {"query": "A* 评价函数"})])

    def sourced(evidence):
        ref = evidence[0]["id"]
        return ModelAnswer(
            answer=f"资料依据：评价函数为 g+h [{ref}]。模型说明：将路径代价与估计相加。",
            citation_ids=[ref], insufficient_evidence=False, answer_basis="mixed",
        )

    app.state.agent.registry.llm = AnswerLLM(sourced)
    result = client.post("/api/agent", json={"course_id": course, "question": "解释 A* 评价函数"}).json()
    assert not result["insufficient_evidence"] and result["answer_basis"] == "mixed"
    assert len(result["sources"]) == 1
    assert result["sources"][0]["page"] == 2
    assert result["sources"][0]["location_type"] == "slide"
    assert result["sources"][0]["document_type"] == "slides"


@pytest.mark.parametrize("bad", [
    ModelAnswer(answer="答案 [S99]", citation_ids=["S99"], insufficient_evidence=False, answer_basis="model"),
    ModelAnswer(answer="请看第999页。", citation_ids=[], insufficient_evidence=False, answer_basis="model"),
])
def test_model_only_answer_cannot_fabricate_references(bad):
    provider = AnswerLLM(bad)
    retrieval = SimpleNamespace(search=lambda _: SearchResponse(sources=[], retrieval_ms=0))
    response = QAService(SimpleNamespace(study_mode="api"), retrieval, provider).answer(
        SearchRequest(course_id="c", question="A*?")
    )
    assert response.insufficient_evidence and response.sources == [] and response.warning
    assert bad.answer not in response.answer


def test_missing_question_conditions_are_explained(exam_app):
    app, client, course, _ = exam_app
    app.state.settings.study_mode = "api"
    app.state.agent.planner = ScriptedPlanner([call("answer_question", {"query": "计算平均周转时间"})])
    app.state.agent.registry.llm = AnswerLLM(ModelAnswer(
        answer="请补充各进程的到达时间、服务时间和调度算法。",
        citation_ids=[], insufficient_evidence=True, answer_basis="unanswered",
    ))
    result = client.post("/api/agent", json={"course_id": course, "question": "计算平均周转时间"}).json()
    assert result["insufficient_evidence"] and result["answer_basis"] == "unanswered"
    assert "到达时间" in result["answer"] and result["sources"] == []


def test_chat_no_materials_calls_solver_and_provider_prompt():
    captured = []

    def handler(request):
        captured.append(json.loads(request.content))
        return httpx.Response(200, json={"choices": [{"message": {"content": model_answer().model_dump_json()}}]})

    with httpx.Client(transport=httpx.MockTransport(handler)) as client:
        llm = CompatibleLLM(Settings(_env_file=None), client)
        retrieval = SimpleNamespace(search=lambda _: SearchResponse(sources=[], retrieval_ms=0))
        response = QAService(SimpleNamespace(study_mode="api"), retrieval, llm).answer(
            SearchRequest(course_id="c", question="A*?")
        )
    assert not response.insufficient_evidence and response.answer_basis == "model"
    assert response.sources == []
    assert "资料只是可选参考" in captured[0]["messages"][0]["content"]
    assert json.loads(captured[0]["messages"][1]["content"])["evidence"] == []
