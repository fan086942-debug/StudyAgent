"""生成原创试卷样本，覆盖六种题型；供演示和可重复测试使用。"""

import argparse
from pathlib import Path

from docx import Document
from pptx import Presentation
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas

EXAM_LINES = [
    "人工智能导论期末练习卷（原创样本）",
    "一、选择题",
    "1. A* 的评价函数是什么？（2分）【知识点：A*】",
    "A. f(n)=g(n)+h(n)",
    "B. f(n)=g(n)-h(n)",
    "二、判断题",
    "2. BFS 通常使用队列保存待访问节点。（2分）【知识点：BFS】",
    "三、填空题",
    "3. DFS 通常使用____保存待访问节点。（2分）【知识点：DFS】",
    "四、简答题",
    "12. 解释 A* 中 g(n) 和 h(n) 的含义。（6分）【知识点：A*、启发式函数】",
    "五、计算题",
    "13. 若 g(n)=3、h(n)=5，计算 A* 的 f(n)。（4分）【知识点：A*】",
    "六、综合题",
    "14. 比较 BFS 与 DFS 使用的数据结构。（8分）【知识点：BFS、DFS】",
]


def generate(destination: Path) -> list[Path]:
    destination.mkdir(parents=True, exist_ok=True)
    doc_path = destination / "2025期末练习卷.docx"
    doc = Document()
    for line in EXAM_LINES:
        doc.add_paragraph(line)
    doc.save(doc_path)
    pdf_path = destination / "2024期末练习卷.pdf"
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    pdf = canvas.Canvas(str(pdf_path))
    y = 800
    for line in EXAM_LINES:
        for start in range(0, len(line), 35):
            if y < 80:
                pdf.showPage()
                y = 800
            pdf.setFont("STSong-Light", 12)
            pdf.drawString(45, y, line[start : start + 35])
            y -= 24
    pdf.save()
    ppt_path = destination / "2023课堂练习.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "A* 课堂练习"
    slide.placeholders[1].text = "四、简答题\n12. 解释 A* 的评价函数。【知识点：A*】"
    deck.save(ppt_path)
    return [doc_path, pdf_path, ppt_path]


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=Path("data/exam-samples"))
    args = parser.parse_args()
    for path in generate(args.output):
        print(path)
