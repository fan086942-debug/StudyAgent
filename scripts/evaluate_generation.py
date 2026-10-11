"""原创小样本的生成流程验收；离线摘录不代表 LLM 答案质量。"""

import argparse
import json
import tempfile
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient

from app.core.config import Settings
from app.main import create_app

PARAGRAPHS = [
    "A* 评价函数为 f(n)=g(n)+h(n)，用于比较待扩展节点。",
    "A* 评价函数中的 g(n) 表示从起点到当前节点的实际路径代价。",
    "A* 评价函数中的 h(n) 表示从当前节点到目标的估计代价。",
    "BFS 使用先进先出的队列，按层次扩展节点。",
    "BFS 在每条边代价相同时可找到边数最少的路径。",
    "DFS 使用后进先出的栈，沿分支深入搜索。",
]


def make_sample(path):
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading("搜索算法 · 原创验收笔记", 0)
    for paragraph in PARAGRAPHS:
        doc.add_paragraph(paragraph)
    doc.save(path)
    return path


def evaluate(mode):
    rows = []
    with tempfile.TemporaryDirectory(prefix="study-generation-", ignore_cleanup_errors=True) as tmp:
        root = Path(tmp)
        settings = Settings(
            _env_file=".env" if mode == "api" else None,
            study_mode=mode,
            data_dir=root / "db",
        )
        with TestClient(create_app(settings)) as client:
            course = client.post("/api/courses", json={"name": "Phase 3 验收"}).json()["id"]
            sample = make_sample(root / "notes.docx")
            response = client.post(
                f"/api/courses/{course}/documents",
                files={"file": (sample.name, sample.read_bytes())},
                data={"document_type": "notes"},
            )
            response.raise_for_status()
            for kind, topics, count, expected in [
                ("review", ["A* 评价函数"], 3, 201),
                ("review", ["A* 评价函数", "BFS 队列"], 3, 201),
                ("mock_exam", ["A* 评价函数"], 3, 201),
                ("review", ["量子纠缠贝尔不等式"], 3, 409),
            ]:
                response = client.post(
                    "/api/generate",
                    json={
                        "course_id": course,
                        "kind": kind,
                        "topics": topics,
                        "question_count": count,
                    },
                )
                data = response.json()
                row = {
                    "kind": kind,
                    "topics": topics,
                    "expected_status": expected,
                    "actual_status": response.status_code,
                    "passed": response.status_code == expected,
                    "manual_factual_correctness": None,
                }
                if response.status_code == 201:
                    url = f"/api/courses/{course}/artifacts/{data['id']}"
                    row.update(
                        source_count=len(data["sources"]),
                        question_count=len(data["content"]["questions"]),
                        content=data["content"],
                        usage=data["usage"],
                        saved=client.get(url).json() == data,
                        download_status=client.get(url + "/download").status_code,
                    )
                    row["passed"] &= row["saved"] and row["download_status"] == 200
                else:
                    row["error"] = data
                rows.append(row)
    return {
        "created_at": datetime.now(timezone.utc).isoformat(),
        "mode": mode,
        "note": "4 个原创小样本，默认检索阈值。demo 仅验证摘录与保存流程；API 结果仍需人工核对。",
        "passed": sum(row["passed"] for row in rows),
        "cases": len(rows),
        "results": rows,
    }


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--mode", choices=["demo", "api"], default="demo")
    parser.add_argument("--output", type=Path, default=Path("data/evaluation-generation-demo.json"))
    args = parser.parse_args()
    report = evaluate(args.mode)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({k: v for k, v in report.items() if k != "results"}, ensure_ascii=False))
