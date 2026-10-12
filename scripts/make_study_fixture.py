"""Create an authored synthetic textbook for isolated study-pack acceptance."""

from pathlib import Path

from docx import Document


def generate(path: Path):
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Document()
    for chapter in range(1, 10):
        book.add_heading(f"第{chapter}章 学习专题{chapter}", level=1)
        for index in range(3):
            book.add_paragraph(f"专题{chapter}要点{index+1}：学习模型应区分训练资料、验证资料和测试资料。训练用于拟合，验证用于调整，测试用于独立评估。")
    book.save(path)


if __name__ == "__main__":
    generate(Path("data/study-test-fixtures/合成机器学习教材.docx"))
