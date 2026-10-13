"""Real model, synthetic three-chapter material, isolated storage; never reads user documents."""

import argparse
import json
from io import BytesIO
from pathlib import Path
from uuid import uuid4

from docx import Document
from fastapi.testclient import TestClient
from pydantic import SecretStr

from app.core.config import Settings
from app.main import create_app
from app.services.connection import model_services


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--topic-only", action="store_true", help="Only verify the two-question topic request"
    )
    parser.add_argument(
        "--memory", action="store_true", help="Also verify topic and question follow-ups"
    )
    args = parser.parse_args()
    profile = json.loads(Path("data/model-connection.json").read_text(encoding="utf-8"))
    settings = Settings(auth_enabled=False,
        _env_file=None, data_dir=Path("data/model-learning-smoke"), min_similarity=0
    )
    app = create_app(settings)
    with TestClient(app) as client:
        configured = settings.model_copy(
            update=dict(
                study_mode="api",
                llm_base_url=profile["base_url"],
                llm_model=profile["model"],
                llm_api_key=SecretStr(profile["api_key"]),
            )
        )
        app.state.connection._bind(model_services(app.state, configured))
        course = client.post("/api/courses", json={"name": "真实模型 · 自编资料验证"}).json()["id"]
        doc = Document()
        for title, text in (
            (
                "第1章 层次聚类",
                "凝聚层次聚类从单点簇开始，反复合并距离最近的两个簇，得到层次结构。单链接用跨簇点对距离的最小值，完全链接用最大值，平均链接用所有跨簇点对距离的平均值。欧氏距离是对应坐标差的平方和再开方。",
            ),
            (
                "第2章 K均值",
                "K均值反复将样本分配给最近的中心，再将各簇中心更新为样本坐标的平均值。目标是最小化样本到各自簇中心的平方欧氏距离之和。需要预先指定簇数，初始化可能影响结果，异常值可能影响均值。",
            ),
            (
                "第3章 泛化与验证",
                "训练资料用于拟合模型，验证资料用于调整超参数，测试资料用于最终评估。过拟合常表现为训练误差低而测试误差高。交叉验证通过多次轮换划分估计泛化表现，测试资料不能用于反复调参。",
            ),
        ):
            doc.add_heading(title, level=1)
            doc.add_paragraph(text)
        buffer = BytesIO()
        doc.save(buffer)
        uploaded = client.post(
            f"/api/courses/{course}/documents",
            files={"file": ("自编三章学习资料.docx", buffer.getvalue())},
            data={"document_type": "textbook"},
        )
        assert uploaded.status_code == 201, uploaded.text
        document = uploaded.json()["id"]
        if args.topic_only:
            identifier = uuid4().hex if args.memory else None
            result = client.post(
                "/api/agent",
                json=dict(
                    course_id=course,
                    document_ids=[document],
                    question="根据层次聚类出两道题目",
                    session_id=identifier,
                ),
            ).json()
            assert result["finish_reason"] == "completed" and result["artifacts"], result["answer"]
            assert [step["tool"] for step in result["trace"]] == ["generate_mock_exam"]
            artifact = client.get(
                f"/api/courses/{course}/artifacts/{result['artifacts'][0]['id']}"
            ).json()
            questions = artifact["content"]["questions"]
            assert len(questions) == 2 and not artifact["content"]["sections"]
            assert artifact["request"]["topics"] == ["层次聚类"]
            assert {source["document_id"] for source in artifact["sources"]} == {document}
            assert all(
                item["answer"] and item["explanation"] and item["citation_ids"]
                for item in questions
            )
            report = dict(
                model=profile["model"],
                material="自编合成资料（不包含用户教材）",
                request="根据层次聚类出两道题目",
                tools=[step["tool"] for step in result["trace"]],
                questions=questions,
                usage=result["usage"],
                success=True,
            )
            Path("docs/evaluation-topic-generation.json").write_text(
                json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
            )
            print(
                json.dumps(
                    {k: v for k, v in report.items() if k != "questions"}, ensure_ascii=False
                ),
                flush=True,
            )
            if args.memory:
                follow = client.post(
                    "/api/agent",
                    json=dict(
                        course_id=course,
                        document_ids=[document],
                        session_id=identifier,
                        question="再来两道题",
                    ),
                ).json()
                assert follow["artifacts"] and follow["context_turns"] == 2, follow["answer"]
                next_artifact = client.get(
                    f"/api/courses/{course}/artifacts/{follow['artifacts'][0]['id']}"
                ).json()
                assert len(next_artifact["content"]["questions"]) == 2
                assert {q["stem"] for q in next_artifact["content"]["questions"]}.isdisjoint(
                    q["stem"] for q in questions
                )
                explain = client.post(
                    "/api/agent",
                    json=dict(
                        course_id=course,
                        document_ids=[document],
                        session_id=identifier,
                        question="解释一下第二题，为什么是这样？",
                    ),
                ).json()
                assert (
                    explain["finish_reason"] == "completed" and not explain["insufficient_evidence"]
                ), explain["answer"]
                assert explain["context_turns"] == 3 and explain["sources"]
                memory_report = dict(
                    material="自编合成资料（不包含用户教材）",
                    model=profile["model"],
                    turns=3,
                    topic_followup=True,
                    generated_question_followup=True,
                    usage=[result["usage"], follow["usage"], explain["usage"]],
                    question=next_artifact["content"]["questions"][1],
                    explanation=explain["answer"],
                    success=True,
                )
                Path("docs/evaluation-agent-memory.json").write_text(
                    json.dumps(memory_report, ensure_ascii=False, indent=2), encoding="utf-8"
                )
                print(
                    json.dumps(
                        {
                            k: v
                            for k, v in memory_report.items()
                            if k not in {"question", "explanation"}
                        },
                        ensure_ascii=False,
                    ),
                    flush=True,
                )
            return
        paper = client.post(
            "/api/agent",
            json=dict(
                course_id=course,
                document_ids=[document],
                question="根据教材每章节重点为我出一套卷子，3道简答题",
            ),
        ).json()
        assert paper["artifacts"], paper["answer"]
        artifact = client.get(
            f"/api/courses/{course}/artifacts/{paper['artifacts'][0]['id']}"
        ).json()
        assert (
            len(artifact["content"]["sections"]) == 3 and len(artifact["content"]["questions"]) == 3
        )
        answer = client.post(
            "/api/agent",
            json=dict(
                course_id=course,
                document_ids=[document],
                question="层次聚类的概念和它如何计算距离？",
            ),
        ).json()
        assert answer["finish_reason"] == "completed" and not answer["insufficient_evidence"], (
            answer["answer"]
        )
        assert answer["sources"]
        report = dict(
            model=profile["model"],
            material="自编合成三章资料（不包含用户教材）",
            chapters=3,
            questions=3,
            study_usage=paper["usage"],
            question_usage=answer["usage"],
            question_tools=[t["tool"] for t in answer["trace"]],
            answer=answer["answer"],
            coverage=artifact["coverage"],
            success=True,
        )
        Path("docs/evaluation-model-learning.json").write_text(
            json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        print(
            json.dumps(
                {k: v for k, v in report.items() if k not in ("answer", "coverage")},
                ensure_ascii=False,
            ),
            flush=True,
        )


if __name__ == "__main__":
    main()
