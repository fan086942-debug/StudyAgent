"""Real API acceptance of natural section exams and local Word export, synthetic inputs only."""

import json
import logging
import tempfile
from io import BytesIO
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient
from pptx import Presentation

from app.core.config import Settings
from app.main import create_app
from app.schemas.connection import ConnectionInput
from app.services.connection import model_services


def main():
    output = Path("data/exam-word-acceptance")
    output.mkdir(parents=True, exist_ok=True)
    profile = ConnectionInput.model_validate_json(Path("data/model-connection.json").read_text("utf-8"))
    report = dict(model=profile.model, synthetic_material=True, cases={})
    deck = Presentation()
    for title, body in [
        ("2.1 进程", "2.1.1 进程状态\n2.1.2 进程调度\n2.1.3 进程通信"),
        ("进程状态与切换", "进程是程序的一次执行过程。就绪态进程已经具备运行条件，等待处理机；"
         "运行态进程正在占用处理机；阻塞态进程正在等待事件。时间片用完时运行态转为就绪态；"
         "请求输入输出并等待时运行态转为阻塞态；事件完成时阻塞态转为就绪态。"),
        ("进程调度", "进程调度从就绪队列选择进程，分配处理机并恢复运行现场。"
         "先来先服务按进入就绪队列的顺序分配；时间片轮转给队首进程分配固定时间片，"
         "未完成的进程返回队尾。短作业优先选择预计运行时间较短者，可能导致长任务等待过久。"),
        ("进程通信与同步", "进程通信可使用共享存储、消息传递和管道。互斥用于避免同时进入临界区，"
         "同步用于满足执行先后约束。临界资源一次只允许一个进程使用。"),
        ("2.2 缓存淘汰", "缓存容量有限，空间不足时需要选择缓存项淘汰。"),
        ("缓存淘汰策略", "先进先出淘汰最早进入缓存的项目；最近最少使用淘汰最长时间未被访问的项目。"
         "命中表示请求数据已经在缓存中，未命中需要从后备存储加载；命中率为命中次数除以总请求次数。"
         "命中率受容量、访问局部性和淘汰策略影响。"),
    ]:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    stream = BytesIO()
    deck.save(stream)
    logging.getLogger("httpx").setLevel(logging.WARNING)
    with tempfile.TemporaryDirectory(prefix="exam-word-", dir=output) as directory:
        app = create_app(Settings(_env_file=None, data_dir=Path(directory), auth_enabled=False, request_timeout=120))
        with TestClient(app) as client:
            app.state.connection._bind(model_services(app.state, app.state.connection._settings(profile)))
            app.state.connection.profile = profile
            course = client.post("/api/courses", json={"name": "出卷与Word验收"}).json()["id"]
            uploaded = client.post(f"/api/courses/{course}/documents", files={
                "file": ("合成进程与缓存课件.pptx", stream.getvalue())}, data={"document_type": "slides"})
            assert uploaded.status_code == 201, uploaded.text
            document = uploaded.json()["id"]
            for topic in ["进程", "缓存淘汰"]:
                question = f"根据这{topic}这部分的知识点出一张试卷"
                result = client.post("/api/agent", json=dict(course_id=course, document_ids=[document], question=question)).json()
                assert result["finish_reason"] == "completed" and result["artifacts"], result
                assert [step["tool"] for step in result["trace"]] == ["generate_mock_exam"]
                identifier = result["artifacts"][0]["id"]
                artifact = client.get(f"/api/courses/{course}/artifacts/{identifier}").json()
                assert artifact["request"]["topics"] == [topic]
                assert len(artifact["content"]["questions"]) == 5
                assert all(q["answer"] and q["explanation"] and q["citation_ids"] for q in artifact["content"]["questions"])
                assert {s["document_id"] for s in artifact["sources"]} == {document}
                for kind in ["qa", "knowledge"]:
                    response = client.get(f"/api/courses/{course}/artifacts/{identifier}/word?kind={kind}")
                    assert response.status_code == 200
                    text = "\n".join(p.text for p in Document(BytesIO(response.content)).paragraphs)
                    assert "资料来源与原文节选" in text and "幻灯片" in text
                    if topic == "进程":
                        (output / f"{kind}.docx").write_bytes(response.content)
                report["cases"][topic] = dict(passed=True, request=question, usage=result["usage"], artifact=artifact)
                print(f"PASS {topic}: five questions, answers, sources, Word downloads", flush=True)
            report["passed"] = True
    (output / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), "utf-8")


if __name__ == "__main__":
    main()
