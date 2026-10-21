"""Real API acceptance of textbook question generation, using isolated storage.

Run: python -m scripts.verify_textbook_generation_api
Uses the saved model connection without printing or persisting its key.
"""

import argparse
import json
import logging
import shutil
import sqlite3
import tempfile
import time
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from uuid import uuid4

from docx import Document
from fastapi.testclient import TestClient

from app.core.config import Settings
from app.main import create_app
from app.models import Document as StoredDocument
from app.schemas.connection import ConnectionInput
from app.services.connection import model_services


def textbook_bytes():
    doc = Document()
    for title, text in (
        ("第1章 层次聚类", "凝聚层次聚类从每个样本独立成簇开始，反复合并距离最近的两个簇。"
         "单链接取跨簇点对距离的最小值，完全链接取最大值，平均链接取平均值。"
         "欧氏距离是各维坐标差的平方和再开方。若两个簇的跨簇点对距离为2、4、6、8，"
         "单链接距离为2，完全链接距离为8，平均链接距离为5。"),
        ("第2章 K均值", "K均值聚类首先指定簇数K并初始化中心，再交替执行样本分配和中心更新。"
         "分配时将样本分给最近的中心，更新时对簇内样本坐标取均值。"
         "目标是最小化簇内平方误差和；异常值及初始化会影响结果。"),
        ("第3章 泛化与验证", "训练集用于拟合模型，验证集用于选择超参数，测试集用于最终评估。"
         "不能反复使用测试集调参。过拟合表现为训练误差低而测试误差高。"
         "交叉验证轮换训练和验证划分，估计模型的泛化表现。"),
    ):
        doc.add_heading(title, level=1)
        doc.add_paragraph(text)
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def require(condition, message):
    if not condition:
        raise AssertionError(message)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=Path("data/textbook-generation-api/report.json"))
    parser.add_argument("--real-textbook", type=Path,
                        default=Path("data/uploads/e5beaaea3b3e42d4805dff7ee38491db.pdf"))
    parser.add_argument("--skip-real-textbook", action="store_true",
                        help="Only use generated synthetic material; do not read the user's PDF")
    parser.add_argument("--only-real-textbook", action="store_true",
                        help="Only run the authorized existing-PDF generation case")
    args = parser.parse_args()
    args.output.parent.mkdir(parents=True, exist_ok=True)
    profile = ConnectionInput.model_validate_json(Path("data/model-connection.json").read_text("utf-8"))
    require(profile.mode == "api" and profile.api_key.get_secret_value(), "Real API is not configured")
    report = dict(started_at=datetime.now(timezone.utc).isoformat(), real_api=True,
                  model=profile.model, cases={}, provider_calls=[],
                  settings=dict(min_similarity=0.25, embedding="local character hash",
                                request_timeout=120, ocr_enabled=True),
                  limitations=["Existing scanned PDF reuses fingerprint-validated local OCR cache; fresh OCR is not tested.",
                               "Whole-book generation tested on a three-chapter synthetic textbook."])
    started = time.perf_counter()

    def save():
        report["elapsed_seconds"] = round(time.perf_counter() - started, 2)
        args.output.write_text(json.dumps(report, ensure_ascii=False, indent=2), "utf-8")

    def case(name, run):
        tick = time.perf_counter()
        try:
            result = run()
            report["cases"][name] = dict(passed=True, seconds=round(time.perf_counter()-tick, 2), **result)
            print(f"PASS {name}", flush=True)
        except Exception as exc:
            report["cases"][name] = dict(passed=False, seconds=round(time.perf_counter()-tick, 2),
                                          error_type=type(exc).__name__, error=str(exc)[:1800])
            print(f"FAIL {name}: {type(exc).__name__}: {str(exc)[:400]}", flush=True)
        save()

    logging.getLogger("httpx").setLevel(logging.WARNING)
    with tempfile.TemporaryDirectory(prefix="textbook-api-", dir="data") as directory:
        settings = Settings(_env_file=None, data_dir=Path(directory), auth_enabled=False,
                            require_real_model=True, request_timeout=120, ocr_enabled=True)
        app = create_app(settings)
        with TestClient(app) as client:
            # Use the real provider and HTTP client; no mock responses or demo generation.
            configured = app.state.connection._settings(profile)
            app.state.connection._bind(model_services(app.state, configured))
            app.state.connection.profile = profile
            app.state.connection._require_model(profile)
            original_post = app.state.http.post

            def observed_post(url, **kwargs):
                tick = time.perf_counter()
                entry = dict(endpoint=url.rsplit("/", 1)[-1], model=kwargs.get("json", {}).get("model"))
                report["provider_calls"].append(entry)
                try:
                    response = original_post(url, **kwargs)
                    entry["status"] = response.status_code
                    if response.status_code == 200:
                        entry["usage"] = response.json().get("usage", {})
                    return response
                finally:
                    entry["seconds"] = round(time.perf_counter()-tick, 2)

            app.state.http.post = observed_post

            def post(url, payload):
                response = client.post(url, json=payload)
                require(response.status_code in {200, 201}, f"HTTP {response.status_code}: {response.text[:1500]}")
                return response.json()

            def create_course(name):
                return post("/api/courses", {"name": name})["id"]

            def upload(course, name, content):
                response = client.post(f"/api/courses/{course}/documents",
                                       files={"file": (name, content)}, data={"document_type": "textbook"})
                require(response.status_code == 201, f"Upload: {response.status_code}: {response.text[:1000]}")
                doc = response.json()
                require(doc["status"] == "ready", f"Document not ready: {doc}")
                return doc

            def validate_artifact(artifact, count, kind, doc_id, course, qtype="short_answer"):
                require(artifact["mode"] == "api", "Unexpected non-API artifact")
                questions = artifact["content"]["questions"]
                require(len(questions) == count, f"Expected {count} questions; got {len(questions)}")
                require({s["document_id"] for s in artifact["sources"]} == {doc_id}, "Wrong source scope")
                require(any(len(s["text"].strip()) > 40 for s in artifact["sources"]),
                        "Evidence contains only headings, without substantive textbook content")
                cites = {s["citation_id"] for s in artifact["sources"]}
                require(all(q["stem"] and q["answer"] and q["explanation"] and q["citation_ids"]
                            and set(q["citation_ids"]) <= cites and q["type"] == qtype for q in questions),
                        "Missing question/answer/explanation, invalid citations or wrong type")
                require(len({q["stem"] for q in questions}) == count, "Duplicate questions")
                base = f"/api/courses/{course}/artifacts/{artifact['id']}"
                require(client.get(base).json() == artifact, "Saved artifact differs")
                answers = client.get(base + "/download")
                paper = client.get(base + "/download?include_answers=false")
                require(answers.status_code == paper.status_code == 200, "Download failed")
                require("参考答案与解析" in answers.text and "参考答案与解析" not in paper.text,
                        "Answer/paper export separation failed")
                (args.output.parent / f"{kind}.md").write_text(answers.text, "utf-8")
                return dict(artifact=artifact, downloads_verified=True, persisted=True)

            if not args.only_real_textbook:
                course = create_course("教材自主出题真实API测试")
                book = upload(course, "三章聚类与验证教材.docx", textbook_bytes())
                document = book["id"]
                report["synthetic_upload"] = book
                report["upload_alone_generated_artifacts"] = len(client.get(f"/api/courses/{course}/artifacts").json())
            session_id = uuid4().hex
            previous_stems = set()

            def agent_questions(name, question, count, qtype="short_answer", memory=False):
                payload = dict(course_id=course, document_ids=[document], question=question)
                if memory:
                    payload["session_id"] = session_id
                result = post("/api/agent", payload)
                require(result["finish_reason"] == "completed" and result["artifacts"],
                        json.dumps(result, ensure_ascii=False)[:1800])
                artifact = client.get(f"/api/courses/{course}/artifacts/{result['artifacts'][0]['id']}").json()
                info = validate_artifact(artifact, count, name, document, course, qtype)
                if name == "whole_textbook_auto_topics":
                    require(len(artifact["content"]["sections"]) == 3,
                            "Whole-material learning did not summarize all three chapters")
                    require(len({q["topic"] for q in artifact["content"]["questions"]}) == 3,
                            "Generated questions did not cover all three chapters")
                info.update(request=question, trace=result["trace"], usage=result["usage"],
                            context_turns=result.get("context_turns"))
                stems = {q["stem"] for q in artifact["content"]["questions"]}
                if memory:
                    require(stems.isdisjoint(previous_stems), "Follow-up repeats earlier questions")
                    previous_stems.update(stems)
                return info

            if not args.only_real_textbook:
                case("whole_textbook_auto_topics", lambda: agent_questions(
                    "whole_textbook_auto_topics", "根据教材每章节重点为我出一套卷子，3道简答题", 3))
                case("topic_short_answer", lambda: agent_questions(
                    "topic_short_answer", "根据层次聚类出两道题目", 2, memory=True))
                case("followup_new_questions", lambda: agent_questions(
                    "followup_new_questions", "再来两道题", 2, memory=True))
                case("topic_choice", lambda: agent_questions(
                    "topic_choice", "根据层次聚类出两道选择题", 2, "choice"))
                case("topic_calculation", lambda: agent_questions(
                    "topic_calculation", "根据层次聚类出一道计算题", 1, "calculation"))

            def no_evidence():
                empty = create_course("无教材出题边界测试")
                before = len(report["provider_calls"])
                result = post("/api/agent", dict(course_id=empty, question="根据层次聚类出两道题目"))
                require(not result["artifacts"], "Generated unsupported artifact")
                require(result["trace"][0]["observation"]["error"]["code"] == "generation_no_evidence",
                        f"Unexpected failure: {result}")
                require(len(report["provider_calls"]) == before, "No-evidence request called the model")
                return dict(response=result, no_provider_call=True)

            if not args.only_real_textbook:
                case("no_evidence_rejected", no_evidence)

            def real_pdf():
                require(args.real_textbook.is_file(), "Existing real textbook PDF missing")
                real_course = create_course("已有机器学习PDF教材出题测试")
                # Recreate an already imported document in the isolated workspace, then
                # run the ordinary retry API. Source PDF/database/cache remain read-only.
                original_id = args.real_textbook.stem
                database = Path("data/studyagent.db").resolve()
                with sqlite3.connect(database.as_uri() + "?mode=ro", uri=True) as source_db:
                    source_db.row_factory = sqlite3.Row
                    row = source_db.execute("SELECT * FROM documents WHERE id=?", (original_id,)).fetchone()
                    require(row is not None, "Existing textbook database record missing")
                    record = dict(row)
                cache = Path("data/ocr") / original_id
                require(cache.is_dir(), "Existing scanned textbook OCR cache missing")
                shutil.copytree(cache, settings.data_dir / "ocr" / original_id)
                record.update(course_id=real_course, storage_path=str(args.real_textbook.resolve()),
                              status="failed", chunk_count=0, embedding_fingerprint=None,
                              error=None, warning=None)
                with app.state.sessions() as db:
                    db.add(StoredDocument(**record))
                    db.commit()
                print("Reindexing existing 442-page textbook with validated OCR cache in isolated workspace...", flush=True)
                real_doc = post(f"/api/documents/{original_id}/retry", None)
                require(real_doc["status"] == "ready", f"Textbook not ready: {real_doc}")
                result = post("/api/agent", dict(course_id=real_course, document_ids=[real_doc["id"]],
                                               question="根据层次聚类出两道题目"))
                require(result["finish_reason"] == "completed" and result["artifacts"],
                        json.dumps(result, ensure_ascii=False)[:1800])
                artifact = client.get(f"/api/courses/{real_course}/artifacts/{result['artifacts'][0]['id']}").json()
                info = validate_artifact(artifact, 2, "real_pdf_topic", real_doc["id"], real_course)
                require(all(s["page"] and s["location_type"] == "pdf_page" for s in artifact["sources"]),
                        "Missing PDF page references")
                info.update(reindexed_document=real_doc, request="根据层次聚类出两道题目", trace=result["trace"],
                            cached_ocr_reused=True)
                return info

            if args.skip_real_textbook:
                report["real_pdf_skipped"] = "No private textbook content read or sent; awaiting explicit authorization."
            else:
                case("real_pdf_topic", real_pdf)
    report["passed"] = all(c["passed"] for c in report["cases"].values())
    report["total_tokens"] = sum(c.get("usage", {}).get("total_tokens", 0) for c in report["provider_calls"])
    save()
    print(json.dumps(dict(passed=report["passed"], cases={k:v["passed"] for k,v in report["cases"].items()},
                          total_tokens=report["total_tokens"], seconds=report["elapsed_seconds"],
                          report=str(args.output)), ensure_ascii=False), flush=True)
    return 0 if report["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
