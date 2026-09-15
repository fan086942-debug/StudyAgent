"""生成自编的小型测试资料；不需要下载或复制真实教材。"""

import argparse
import json
from pathlib import Path

from docx import Document
from pptx import Presentation
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas


def generate(destination: Path) -> list[Path]:
    destination.mkdir(parents=True, exist_ok=True)
    corpus_path = Path(__file__).resolve().parents[1] / "tests/fixtures/corpus.json"
    corpus = json.loads(corpus_path.read_text(encoding="utf-8"))
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    paths = []
    for spec in corpus["documents"]:
        path = destination / spec["name"]
        if path.suffix == ".pdf":
            pdf = canvas.Canvas(str(path))
            for page in spec["pages"]:
                pdf.setFont("STSong-Light", 16)
                pdf.drawString(50, 790, page["title"])
                pdf.setFont("STSong-Light", 12)
                for index in range(0, len(page["text"]), 35):
                    pdf.drawString(50, 745 - (index // 35) * 24, page["text"][index : index + 35])
                pdf.showPage()
            pdf.save()
        elif path.suffix == ".pptx":
            deck = Presentation()
            for page in spec["pages"]:
                slide = deck.slides.add_slide(deck.slide_layouts[1])
                slide.shapes.title.text = page["title"]
                slide.placeholders[1].text = page["text"]
                if page.get("notes"):
                    slide.notes_slide.notes_text_frame.text = page["notes"]
            deck.save(path)
        else:
            doc = Document()
            for page in spec["pages"]:
                doc.add_heading(page["title"], level=1)
                doc.add_paragraph(page["text"])
            table = doc.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "算法"
            table.cell(0, 1).text = "数据结构"
            table.cell(1, 0).text = "BFS"
            table.cell(1, 1).text = "队列"
            doc.save(path)
        paths.append(path)
    return paths


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=Path("data/samples"))
    args = parser.parse_args()
    for generated in generate(args.output):
        print(generated)
