"""Create local Word copies from validated learning results, without another LLM call."""

import re
from collections import OrderedDict
from io import BytesIO

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def safe_text(value):
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(value or ""))


def paragraph(document, text, style=None):
    node = document.add_paragraph(style=style)
    node.paragraph_format.keep_together = True
    for part in re.split(r"(\*\*.*?\*\*|`[^`]+`)", safe_text(text)):
        if part.startswith("**") and part.endswith("**"):
            node.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            node.add_run(part[1:-1]).font.name = "Consolas"
        else:
            node.add_run(part)
    return node


def markdown(document, text):
    in_code = False
    for line in safe_text(text).splitlines():
        if line.strip().startswith("```"):
            in_code = not in_code
            continue
        if not line.strip():
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)", line)
        bullet = re.match(r"^\s*[-*+]\s+(.+)", line)
        if heading and not in_code:
            document.add_heading(heading[2], min(3, len(heading[1]) + 1))
        elif bullet and not in_code:
            paragraph(document, bullet[1], "List Bullet")
        else:
            node = paragraph(document, line)
            if in_code:
                for run in node.runs:
                    run.font.name = "Consolas"


def new_document(title, introduction):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.2)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "List Bullet"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        properties = style.element.find(qn("w:pPr"))
        if properties is not None:
            for border in list(properties.findall(qn("w:pBdr"))):
                properties.remove(border)
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.3
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    doc.styles["Title"].font.size = Pt(20)
    doc.styles["Title"].font.color.rgb = RGBColor.from_string("253858")
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        doc.styles[name].paragraph_format.keep_with_next = True
        doc.styles[name].paragraph_format.space_before = Pt(12)
    doc.add_heading(safe_text(title), 0)
    paragraph(doc, introduction)
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run("StudyAgent · ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    return doc


def references(document, ids):
    if ids:
        paragraph(document, "资料依据：" + " ".join(f"[{item}]" for item in ids))


def sources_section(document, sources):
    if not sources:
        return
    document.add_heading("资料来源与原文节选", 1)
    for source in sources:
        location = (
            f"PDF 物理第 {source['page']} 页" if source.get("file_type") == "pdf" and source.get("page")
            else f"第 {source['page']} 张幻灯片" if source.get("location_type") == "slide"
            else f"第 {source.get('block_index', 0)} 个段落或表格块"
        )
        document.add_heading(safe_text(f"[{source['citation_id']}] {source['document_name']} · {location}"), 2)
        text = safe_text(source.get("text"))
        paragraph(document, text[:1600] + ("\n（原文节选至此）" if len(text) > 1600 else ""))


def finish(document):
    result = BytesIO()
    document.save(result)
    return result.getvalue()


def artifact_word(artifact, kind):
    data = artifact.model_dump(mode="json")
    content = data["content"]
    label = "题目与参考答案" if kind == "qa" else "知识重点整理"
    doc = new_document(content["title"] + " · " + label,
                       "本文件整理本次学习结果，供本地保存与复习；参考答案请结合资料核对。")
    if kind == "qa" and content["questions"]:
        for index, question in enumerate(content["questions"], 1):
            first = len(doc.paragraphs)
            doc.add_heading(f"第 {index} 题 · {question['topic']}", 1)
            markdown(doc, question["stem"])
            for letter, option in zip("ABCD", question["options"]):
                paragraph(doc, f"{letter}. {option}")
            doc.add_heading("参考答案", 2)
            markdown(doc, question["answer"])
            doc.add_heading("解析", 2)
            markdown(doc, question["explanation"])
            references(doc, question["citation_ids"])
            for node in doc.paragraphs[first:-1]:
                node.paragraph_format.keep_with_next = True
    else:
        for section in content["sections"]:
            doc.add_heading(safe_text(section["topic"]), 1)
            markdown(doc, section["summary"])
            doc.add_heading("知识重点", 2)
            for point in section["key_points"]:
                paragraph(doc, point, "List Bullet")
            if section["pitfalls"]:
                doc.add_heading("易错点", 2)
                for point in section["pitfalls"]:
                    paragraph(doc, point, "List Bullet")
            references(doc, section["citation_ids"])
        if not content["sections"]:
            topics = OrderedDict()
            for question in content["questions"]:
                topics.setdefault(question["topic"], []).append(question)
            for topic, questions in topics.items():
                doc.add_heading(safe_text(topic), 1)
                paragraph(doc, "以下知识要点来自本次练习的参考答案与解析。")
                for question in questions:
                    markdown(doc, question["answer"])
                    markdown(doc, question["explanation"])
                    references(doc, question["citation_ids"])
    if data["warning"]:
        paragraph(doc, data["warning"])
    sources_section(doc, data["sources"])
    return finish(doc)


def conversation_word(record, kind):
    label = "题目与参考答案" if kind == "qa" else "知识重点整理"
    doc = new_document(label, "本文件保留本次提问、已有解答和引用资料，供本地保存与复习。")
    basis = record["result"].get("answer_basis")
    paragraph(doc, {
        "model": "解答依据：模型推理，本次未引用教材或课件。",
        "mixed": "解答依据：资料与模型推理，请结合来源原文核对。",
        "materials": "解答依据：课程资料，引用原文及位置见文末。",
    }.get(basis, "解答依据请结合本次引用核对。"))
    doc.add_heading("本次提问" if kind == "qa" else "学习主题", 1)
    paragraph(doc, record["question"])
    doc.add_heading("参考答案与解析" if kind == "qa" else "知识讲解与要点", 1)
    markdown(doc, record["result"]["answer"])
    if record["result"].get("warning"):
        paragraph(doc, record["result"]["warning"])
    sources_section(doc, record["result"].get("sources", []))
    return finish(doc)
