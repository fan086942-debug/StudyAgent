from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.schemas.document import ParsedDocument, TextBlock


def shape_text(shapes) -> list[str]:
    parts = []
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            parts.extend(shape_text(shape.shapes))
        elif shape.has_table:
            parts.extend(" | ".join(cell.text for cell in row.cells) for row in shape.table.rows)
        elif shape.has_text_frame:
            parts.append(shape.text_frame.text)
    return parts


def parse_pptx(path: Path) -> ParsedDocument:
    presentation = Presentation(path)
    blocks = []
    for number, slide in enumerate(presentation.slides, 1):
        title = slide.shapes.title.text if slide.shapes.title is not None else None
        parts = shape_text(slide.shapes)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes is not None and notes.text.strip():
                parts.append("备注：" + notes.text)
        blocks.append(
            TextBlock(
                text="\n".join(parts),
                page=number,
                section=title,
                block_index=number,
                location_type="slide",
            )
        )
    return ParsedDocument(blocks=blocks)
